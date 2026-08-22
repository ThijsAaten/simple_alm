"""
Convert the calibration source data to CSV.

WHY THIS EXISTS. The equity factor calibration originally read two pandas pickles.
Pickles are version-coupled: these were written by numpy 2.x and CANNOT be read
by the environment this repository declares in requirements.txt (numpy>=1.26) —
`read_pickle` fails with `ModuleNotFoundError: numpy._core.numeric`. Committing
them would have produced the appearance of reproducibility without the substance.
Pickle is also opaque to diff, unreviewable, and executes arbitrary code on load.

CSV costs ~140KB for the whole dataset, reproduces the published calibration table
to its full stated precision, and can be read by anything.

TRUNCATION. The Bloomberg workbook was pulled with an end date of 2026-12-31,
which is in the future. `Fill=P` carries the last observation forward, so every
leg repeats its 2026-07 value through December. This script truncates at
2026-07-31 for that reason — the same truncation the FX loading derivation
applied. Re-running without it would silently dilute every estimate with a run of
zero returns.

Run from the repo root (needs numpy>=2 to read the legacy pickles ONE time; after
that the CSVs are the source and plain numpy 1.x is enough):

    python calibration/convert_source_data.py
"""
from pathlib import Path

import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
CUTOFF = "2026-07-31"          # last month with real observations; see TRUNCATION above
FLOAT_FMT = "%.17g"            # round-trips to within 1 ulp (~1e-16), far below data precision


def convert_pickles() -> None:
    """ret_usd.pkl and fx_levels.pkl -> CSV, unchanged in content or convention."""
    for stem in ("ret_usd", "fx_levels"):
        src = DATA / f"{stem}.pkl"
        if not src.exists():
            print(f"  {stem}.pkl absent — skipping (CSV is the source now)")
            continue
        df = pd.read_pickle(src)
        df.index.name = "date"
        out = DATA / f"{stem}.csv"
        df.to_csv(out, float_format=FLOAT_FMT)
        print(f"  {stem}.pkl -> {out.name}  ({df.shape[0]} rows x {df.shape[1]} cols)")


def convert_workbook() -> None:
    """data/fx_history_bloomberg.xlsx -> fx_bloomberg_legs.csv.

    The workbook lays each currency out as an independent Date|Value block on a
    three-column pitch. Blocks are read by their row-3 header and joined on date.
    """
    wb = openpyxl.load_workbook(DATA / "fx_history_bloomberg.xlsx", data_only=True)
    ws = wb["Data"]
    series = {}
    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=3, column=col).value
        if not header:
            continue
        name = str(header).split(" (")[0].strip()      # drop "(optional)" etc.
        rows = {}
        for r in range(5, ws.max_row + 1):
            d = ws.cell(row=r, column=col).value
            v = ws.cell(row=r, column=col + 1).value
            if d is None or v is None:
                break
            rows[pd.Timestamp(d)] = float(v)
        series[name] = pd.Series(rows)
    df = pd.DataFrame(series).sort_index()
    full = len(df)
    df = df.loc[:CUTOFF]
    df.index.name = "date"
    out = DATA / "fx_bloomberg_legs.csv"
    df.to_csv(out, float_format=FLOAT_FMT)
    print(f"  fx_history_bloomberg.xlsx -> {out.name}  ({df.shape[0]} rows x {df.shape[1]} cols; "
          f"{full - len(df)} carried-forward rows after {CUTOFF} dropped)")
    print(f"     legs: {', '.join(df.columns)}")


if __name__ == "__main__":
    print("Converting calibration source data to CSV:")
    convert_pickles()
    convert_workbook()
